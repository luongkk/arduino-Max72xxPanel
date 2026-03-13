#!/usr/bin/env python3
"""Word -> PowerPoint automation starter CLI.

This is a practical fallback renderer for quickly generating editable PPTX files.
It supports:
- Parsing headings/bullets from a .docx report.
- Optional image bindings by slide index.
- Optional animation specs (stored in slide notes for native runners).
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


def _lazy_imports() -> tuple[Any, Any]:
    try:
        import docx  # type: ignore
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise SystemExit(
            "Missing dependencies. Install with: pip install python-docx python-pptx"
        ) from exc
    return docx, Presentation


@dataclass
class SlidePlan:
    title: str
    bullets: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def parse_docx_to_plan(docx_path: Path) -> list[SlidePlan]:
    docx, _ = _lazy_imports()
    document = docx.Document(str(docx_path))

    slides: list[SlidePlan] = []
    current: SlidePlan | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = (paragraph.style.name or "").lower()
        if "heading 1" in style_name or "heading 2" in style_name:
            if current:
                slides.append(current)
            current = SlidePlan(title=text)
            continue

        if current is None:
            current = SlidePlan(title="Tổng quan")

        if style_name.startswith("list") or text.startswith(("- ", "• ")):
            current.bullets.append(text.lstrip("-• ").strip())
        else:
            current.bullets.append(text)

    if current:
        slides.append(current)

    return slides


def apply_spec(slides: list[SlidePlan], spec: dict[str, Any], base_dir: Path) -> tuple[dict[int, str], dict[int, list[dict[str, Any]]]]:
    image_map: dict[int, str] = {}
    anim_map: dict[int, list[dict[str, Any]]] = {}

    for binding in spec.get("asset_bindings", []):
        slide_idx = int(binding.get("slide_index", 0))
        source = binding.get("source")
        if source:
            image_map[slide_idx] = str((base_dir / source).resolve())

    for item in spec.get("animations", []):
        slide_idx = int(item.get("slide_index", 0))
        timeline = item.get("timeline", [])
        if timeline:
            anim_map[slide_idx] = timeline

    return image_map, anim_map


def render_ppt(
    slides: list[SlidePlan],
    output_pptx: Path,
    template_pptx: Path | None,
    image_map: dict[int, str],
    anim_map: dict[int, list[dict[str, Any]]],
) -> None:
    _, Presentation = _lazy_imports()
    prs = Presentation(str(template_pptx)) if template_pptx else Presentation()

    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]

    for i, plan in enumerate(slides, start=1):
        layout = title_layout if i == 1 else content_layout
        slide = prs.slides.add_slide(layout)

        if slide.shapes.title:
            slide.shapes.title.text = plan.title

        body_placeholder = None
        for ph in slide.placeholders:
            if ph.placeholder_format.idx != 0:
                body_placeholder = ph
                break

        if body_placeholder and hasattr(body_placeholder, "text_frame"):
            tf = body_placeholder.text_frame
            tf.clear()
            for idx, bullet in enumerate(plan.bullets):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = bullet
                p.level = 0

        if i in image_map:
            image_path = Path(image_map[i])
            if image_path.exists():
                slide.shapes.add_picture(str(image_path), left=prs.slide_width * 0.62, top=prs.slide_height * 0.2, width=prs.slide_width * 0.32)
                plan.notes.append(f"Image bound: {image_path.name}")
            else:
                plan.notes.append(f"Image missing: {image_path}")

        if i in anim_map:
            plan.notes.append("Animation timeline (for native PowerPoint runner):")
            plan.notes.append(json.dumps(anim_map[i], ensure_ascii=False))

        if plan.notes:
            notes = slide.notes_slide.notes_text_frame
            notes.text = "\n".join(plan.notes)

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate PPTX from a structured Word report.")
    parser.add_argument("--docx", required=True, type=Path, help="Path to source .docx file")
    parser.add_argument("--out", required=True, type=Path, help="Path to output .pptx file")
    parser.add_argument("--template", type=Path, help="Optional .pptx template")
    parser.add_argument("--spec", type=Path, help="Optional JSON for image/animation instructions")
    args = parser.parse_args()

    slides = parse_docx_to_plan(args.docx)

    spec = {}
    if args.spec:
        spec = json.loads(args.spec.read_text(encoding="utf-8"))

    image_map, anim_map = apply_spec(slides, spec, base_dir=args.spec.parent if args.spec else Path.cwd())
    render_ppt(slides, args.out, args.template, image_map, anim_map)
    print(f"Generated: {args.out}")
    print("Note: Full animation fidelity needs native PowerPoint execution layer (Office JS/COM).")


if __name__ == "__main__":
    main()
